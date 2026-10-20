#!/usr/bin/env python3
"""Build the CS324 Assignment 2 report.

The Word document is intentionally plain OOXML generated with python-docx: no
pandoc/LibreOffice dependency, and it can be regenerated after the group fills
in the highlighted placeholders. build_pdf.ps1 exports the Word file to the
submission PDF with Microsoft Word.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "report" / "CS324_A2_Report.docx"
GIT_LOG = ROOT / "report" / "git-log.txt"

NAVY = RGBColor(0x1F, 0x38, 0x64)
STEEL = RGBColor(0x2E, 0x5A, 0x88)
GREY = RGBColor(0x59, 0x59, 0x59)

ELECTION_PROTO = """syntax = "proto3";
package distrilab;
option java_package = "com.cs324.election.proto";
option java_multiple_files = true;

service ElectionService {
  rpc SendElection(ElectionMessage) returns (ElectionAck);
  rpc SendEcho(EchoMessage) returns (EchoAck);
  rpc AnnounceCoordinator(CoordinatorAnnouncement) returns (CoordinatorAck);
}

message ElectionMessage {
  int32 origin_id = 1;
  int32 candidate_id = 2;
  int32 candidate_jac = 3;
  int64 message_id = 4;
  int64 term = 5;
  repeated int32 voters = 6;
}
message ElectionAck {}

message EchoMessage {
  int64 message_id = 1;
  int32 from_worker_id = 2;
  int32 best_candidate_id = 3;
  int32 best_candidate_jac = 4;
  int64 term = 5;
  repeated int32 voters = 6;
}
message EchoAck {}

message CoordinatorAnnouncement {
  int32 coordinator_id = 1;
  int64 message_id = 2;
  int64 term = 3;
}
message CoordinatorAck {}"""


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 10.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"


def add_table(doc, headers, rows, widths=None, font_size: float = 10.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0].cells
    for index, text in enumerate(headers):
        set_cell_text(header[index], text, bold=True, size=font_size)
        shade(header[index], "D9E2F3")
    for row in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            set_cell_text(cells[index], str(text), size=font_size)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def body(doc, text: str, space_after: float = 6):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return paragraph


def lead(doc, heading: str, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(heading + " ")
    run.bold = True
    run.font.color.rgb = STEEL
    paragraph.add_run(text)
    return paragraph


def bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(item, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)


def numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(item, style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)


def reference(doc, text: str):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.left_indent = Inches(0.35)
    paragraph.paragraph_format.first_line_indent = Inches(-0.35)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def code_block(doc, text: str, size: float = 8):
    for line in text.splitlines() or [""]:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(size)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def page_number_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for label in ("CS324 Assignment 2  |  page ", None, " of ", None):
        if label is None:
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE" if not paragraph.runs or "page" in paragraph.text else "NUMPAGES")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "1"
            run.append(text)
            field.append(run)
            paragraph._p.append(field)
        else:
            run = paragraph.add_run(label)
            run.font.size = Pt(9)
            run.font.color.rgb = GREY


def add_field(paragraph, instruction: str) -> None:
    begin_run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(fld_begin)

    instruction_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    instruction_run._r.append(instr)

    separate_run = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(fld_sep)

    placeholder_run = paragraph.add_run("1")

    end_run = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    end_run._r.append(fld_end)

    # Keep the generated field text in the document footer's run formatting.
    placeholder_run.font.size = Pt(9)
    placeholder_run.font.color.rgb = GREY


def footer_fields(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("CS324 Assignment 2  |  page ")
    add_field(paragraph, "PAGE")
    paragraph.add_run(" of ")
    add_field(paragraph, "NUMPAGES")
    for run in paragraph.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = GREY


def highlight_placeholders(doc) -> None:
    def visit(paragraph):
        for run in paragraph.runs:
            text = run.text
            if "[" in text and "]" in text:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                run.bold = True

    for paragraph in doc.paragraphs:
        visit(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    visit(paragraph)


def heading(doc, text: str, level: int = 1):
    return doc.add_heading(text, level=level)


def page_break(doc):
    doc.add_page_break()


def main() -> None:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)
    footer_fields(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color in (
        ("Heading 1", 23, NAVY),
        ("Heading 2", 15, STEEL),
        ("Heading 3", 12.5, STEEL),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)

    # ------------------------------------------------------------------ title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(70)
    run = title.add_run("CS324 Distributed Systems")
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = STEEL

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Assignment 2")
    run.font.size = Pt(38)
    run.font.bold = True
    run.font.color.rgb = NAVY

    part = doc.add_paragraph()
    part.alignment = WD_ALIGN_PARAGRAPH.CENTER
    part.paragraph_format.space_after = Pt(8)
    run = part.add_run("Part 1: Report")
    run.font.size = Pt(17)
    run.font.bold = True
    run.font.color.rgb = STEEL

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    run = subtitle.add_run("gRPC worker cluster with leader election,\nfailure recovery and streamed progress")
    run.font.size = Pt(16)
    run.font.italic = True
    run.font.color.rgb = GREY

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = ("Course", "Group", "Repository", "Date")
    values = (
        "CS324 Distributed Systems",
        "[Group number / tutorial group]",
        "gRPC-Leader-Election",
        "20 October 2026",
    )
    for index, (label, value) in enumerate(zip(labels, values)):
        set_cell_text(meta.rows[index].cells[0], label, bold=True)
        set_cell_text(meta.rows[index].cells[1], value)
    doc.add_paragraph()

    body(doc, "This report describes the design, implementation and verification of a "
              "distributed worker cluster in which every inter-process operation uses gRPC and "
              "Protocol Buffers. It covers the election protocol, the reasons for choosing gRPC "
              "over Java RMI, failure detection, coordinator state recovery, and the alternatives "
              "considered. The system preserves the Assignment 1 worker behaviours while adding "
              "real-time progress streaming and both worker and coordinator failure tolerance.")
    body(doc, "Group member names, student IDs and contribution percentages are highlighted in "
              "yellow. Replace those placeholders before submission; the percentages must sum to "
              "100%. Meeting-minute details are also highlighted in Appendix C.", space_after=4)

    heading(doc, "Group members and contributions", level=2)
    add_table(
        doc,
        ["Member", "Student ID", "Main responsibilities", "Contribution"],
        [
            ("[Member 1 Name]", "[Student ID]", "Election, quorum, integration, report", "[25%]"),
            ("[Member 2 Name]", "[Student ID]", "Coordinator failover, checkpoint replication, recovery", "[25%]"),
            ("[Member 3 Name]", "[Student ID]", "Job distribution, computation, progress streaming, client UI", "[25%]"),
            ("[Member 4 Name]", "[Student ID]", "Bootstrap, gRPC contracts, tests, build/CI, documentation", "[25%]"),
            ("Total", "", "", "100%"),
        ],
        widths=[1.25, 1.05, 3.0, 0.85],
        font_size=10.5,
    )

    # ------------------------------------------------------------------ Q1
    page_break(doc)
    heading(doc, "Q1. Election protocol and its gRPC interface", level=1)
    body(doc, "The cluster elects one coordinator without a fixed leader or a separate election "
              "server. Each worker is both a client and a server of ElectionService, defined in "
              "shared/src/main/proto/election.proto. The service deliberately contains only three "
              "RPCs: SendElection carries one wave hop, SendEcho carries one subtree answer back "
              "toward the origin, and AnnounceCoordinator publishes the winner. Every call returns "
              "a small acknowledgement; a failed hop is a signal to prune that edge and repair the "
              "topology rather than a reason to retry forever.")

    heading(doc, "Service contract", level=2)
    body(doc, "The frozen contract below is the only interface used for election traffic. Generated "
              "Java stubs are shared by bootstrap, worker, client and integration tests.")
    code_block(doc, ELECTION_PROTO, size=7.5)
    body(doc, "message_id is generated from the sender identity and a local counter, so two roots "
              "cannot accidentally create the same wave. term is the leadership generation. voters "
              "is the union of workers that participated, and is used for the majority check before "
              "an announcement is allowed.")

    heading(doc, "Method responsibilities", level=2)
    lead(doc, "SendElection.", "The root votes for itself and sends its candidate to every neighbour. "
         "A receiver compares (JAC, worker ID), forwards the winning candidate to every neighbour "
         "except its parent, and records the voters it has seen. A duplicate message_id is not "
         "silently dropped: the receiver echoes the best candidate it already knows, so a cycle "
         "cannot leave a parent waiting forever.")
    lead(doc, "SendEcho.", "When all pending children have answered, a node merges their best "
         "candidates and voter sets with its own and sends one echo to its parent. A per-wave lock "
         "prevents a duplicate echo from completing the same wave twice. The root therefore learns "
         "the winner and the full participating set without all-to-all communication.")
    lead(doc, "AnnounceCoordinator.", "The root announces only after the voter union is a strict "
         "majority of ready workers known through the bootstrap directory. Otherwise it records "
         "ELECTION_NO_QUORUM and announces nobody. The announcement carries a newer term; a worker "
         "accepts the (term, coordinator ID) pair once and rejects stale or termless announcements.")

    heading(doc, "Election rule", level=2)
    bullets(doc, [
        "A lower job-allocation counter (JAC) wins; when JACs are equal, the higher worker ID wins. "
        "This preserves the Assignment 1 load-balancing metric instead of electing by ID alone.",
        "A root that cannot collect a strict majority never announces a coordinator, so a minority "
        "partition cannot install a leader by itself.",
        "After the winner is known, QuorumMonitor probes ready workers every second with a 600 ms "
        "deadline, renews a 2.5-second lease, and forces a step-down after two failed checks. A "
        "leader without a majority cannot accept new work, preventing split brain.",
        "Terms increase with each election and are attached to announcements, so a deposed "
        "coordinator cannot reclaim leadership with an old message.",
    ])

    heading(doc, "Why an echo wave", level=2)
    body(doc, "The unstructured network may be a chain, a triangle or a larger mesh, and edges can "
              "disappear while a wave is in flight. The echo wave converges on any connected graph, "
              "uses only local knowledge, and needs no global clock or ring order. It also lets every "
              "node contribute its JAC. The cost is O(E) messages per wave, acceptable at assignment "
              "scale, and the same exchange supplies the voter set used for the majority check.")

    # ------------------------------------------------------------------ Q2
    page_break(doc)
    heading(doc, "Q2. gRPC versus Java RMI", level=1)
    body(doc, "The implementation uses gRPC for every distributed operation: registry discovery, "
              "worker control, election, job submission and assignment, checkpoint replication and "
              "progress streaming. No production class imports java.rmi, extends Remote or exports "
              "an RMI object, and NoRmiDependencyTest enforces that property. The comparison below "
              "is therefore not theoretical: it explains the transport actually used.")

    add_table(
        doc,
        ["Concern", "gRPC + Protocol Buffers", "Java RMI"],
        [
            ("Contract", "Language-neutral .proto schema; generated stubs for many languages.",
             "Java interfaces; Java-only clients."),
            ("Serialization", "Explicit protobuf fields with stable field numbers; inspectable payloads.",
             "Java object serialization; brittle across class changes and harder to inspect."),
            ("Streaming", "First-class unary, client-streaming, server-streaming and bidirectional RPCs.",
             "No equivalent streaming abstraction; callbacks or polling are ad hoc."),
            ("Failure semantics", "Deadlines, cancellation and typed status codes (UNAVAILABLE, DEADLINE_EXCEEDED, ...).",
             "RemoteException and timeouts; less uniform cross-language semantics."),
            ("Transport", "HTTP/2 multiplexing over one connection; flow control and backpressure.",
             "JRMP, a Java-specific protocol; often one request per connection."),
            ("Security", "TLS and token/metadata interceptors are standard; plaintext is an explicit option.",
             "Security requires separate configuration and is commonly omitted in demos."),
            ("Discovery", "Bootstrap RegistryService returns typed WorkerEndpoint messages.",
             "RMI registry or custom lookup; registry is another Java-specific component."),
            ("Build cost", "protoc + gRPC code generation adds a build step.",
             "No code generation, but no schema contract either."),
        ],
        widths=[1.2, 3.0, 2.95],
        font_size=9.5,
    )

    heading(doc, "Why gRPC is the better fit here", level=2)
    body(doc, "First, the assignment requires progress to be pushed to the client as work actually "
              "completes. JobProgressService.WatchJob is a server stream of typed JobProgress "
              "messages, and the client can reconnect with the last known job ID. Re implementing "
              "that in RMI would mean a callback object, manual listener registration and custom "
              "failure handling. Second, the cluster needs explicit per-call deadlines: a dead "
              "worker must not stall a heartbeat, a quorum probe or a task assignment. gRPC "
              "deadlines and status codes make those decisions local and testable. Third, the six "
              "proto files form a frozen, reviewable contract shared by bootstrap, worker, client "
              "and integration tests. That is valuable in a team project because a field-number "
              "change is visible in review, while a Java serialization change can silently break "
              "an older process.")
    body(doc, "The cost is real: the Maven build must run protoc and the gRPC plugin, and generated "
              "sources have to be on the IDE classpath. The project handles this in the shared "
              "module and documents the module classpath in the README. The benefit is a transport "
              "with streaming, deadlines, flow control and a language-neutral contract, which maps "
              "directly to the Assignment 2 requirements. RMI would have been shorter to start but "
              "would have made streaming and cross-process failure handling more custom and more "
              "Java-specific.")

    # ------------------------------------------------------------------ Q3
    page_break(doc)
    heading(doc, "Q3. Detecting failed workers and failed coordinators", level=1)
    body(doc, "Failure detection is deliberately conservative: a slow reply is not treated as a "
              "crash, but prolonged silence eventually removes the process from the set of voters. "
              "The same principle is used at two levels: the coordinator monitors its own majority, "
              "and every worker monitors the coordinator it currently believes in.")

    heading(doc, "Coordinator health", level=2)
    body(doc, "CoordinatorHealthMonitor sends a bounded FailoverService.Heartbeat to the current "
              "coordinator. CoordinatorHealth classifies each result without performing I/O:")
    add_table(
        doc,
        ["State", "Condition", "Response"],
        [
            ("ALIVE", "Successful heartbeat within the 400 ms slow threshold.",
             "Reset the miss counter; keep serving and monitoring."),
            ("SUSPECT", "Successful but slow heartbeat, or one or two consecutive misses.",
             "Keep monitoring; a slow peer is demonstrably up, so no election starts."),
            ("FAILED", "Three consecutive misses, or no success for eight seconds.",
             "FailureGate admits one failover request for this failure window, then survivors elect."),
        ],
        widths=[0.8, 2.85, 3.5],
        font_size=9.5,
    )
    body(doc, "The per-call deadline is 600 ms and the interval is 1000 ms, so a probe cannot overlap "
              "the next one indefinitely. A timeout is only one miss; the threshold of three tolerates "
              "a GC pause or a transient scheduler delay. The eight-second silence rule prevents an "
              "earlier success from masking a peer that has stopped answering mid-window. These "
              "constants are in Constants.java and are covered by CoordinatorHealthTest and "
              "CoordinatorFailoverTest.")
    body(doc, "A coordinator also has to prove that it still has a majority. QuorumMonitor probes "
              "each ready worker's WorkerControlService.GetStatus every second with a 600 ms "
              "deadline. Self and each reachable member count as a vote; the threshold is "
              "floor(n/2)+1. A successful check renews a 2.5-second lease. Two consecutive failed "
              "checks clear the coordinator view and record QUORUM_LOST and "
              "LEADERSHIP_STEP_DOWN. New admissions require a fresh lease, so a leader that has "
              "just lost contact cannot keep serving while a replacement is elected.")

    heading(doc, "Worker failure during a job", level=2)
    body(doc, "The coordinator probes outstanding assignments every 500 ms with a one-second "
              "timeout. If a worker is unreachable, its unfinished slices are reassigned through "
              "fresh discovery, including to a late-joining worker. Each slice has a stable task "
              "ID and an attempt number; the assignment key is (job ID, task ID, attempt). A stale "
              "result from the old owner cannot complete a replacement attempt, and a completed "
              "slice is retained and included once in the aggregate. A failed worker can therefore "
              "lose work without failing the job. The progress stream preserves units that the "
              "worker had already completed and changes only that worker's status to Failed.")

    heading(doc, "Slow versus failed", level=2)
    body(doc, "A successful but slow heartbeat is a SUSPECT signal, not a failure. The distinction "
              "matters because a false failure triggers an unnecessary election and can split work. "
              "The failure threshold and the quorum lease are intentionally longer than one RPC "
              "deadline, but short enough to demonstrate recovery in seconds. A genuinely crashed "
              "process stops answering entirely and eventually crosses both the miss threshold and "
              "the silence bound. TCP errors are treated as failures immediately; a timeout is "
              "treated as a miss until the threshold is reached. The tests exercise both slow and "
              "dead coordinators, minority partitions and leader step-down.")

    # ------------------------------------------------------------------ Q4
    page_break(doc)
    heading(doc, "Q4. Coordinator failure and state recovery", level=1)
    body(doc, "A coordinator is not allowed to acknowledge a job until at least one peer has "
              "confirmed the acceptance snapshot. The snapshot contains the original "
              "SubmitJobRequest, the client callback endpoint, the acceptance time, and every slice "
              "with its input, owner and attempt. The coordinator sends it to published peers over "
              "FailoverService.Replicate and waits up to one second for an acknowledgement whose "
              "sequence covers the snapshot. No peer, a timeout or an RPC error is not confirmation: "
              "the submission returns retryable UNAVAILABLE, no computation starts, and the reserved "
              "job-budget slot is released. At least two live workers are therefore needed to accept "
              "new work.")

    heading(doc, "State recovered after a takeover", level=2)
    add_table(
        doc,
        ["State", "What the checkpoint keeps", "Why it matters"],
        [
            ("Job record", "Job ID, computation type, original request and client endpoint.",
             "Rebuild the submission without asking the client to resend it."),
            ("Task slices", "Task ID, payload, assigned worker and attempt number.",
             "Preserve completed work and identify what must be reassigned."),
            ("Execution result", "Per-slice result and the attempt that produced it.",
             "Retain useful work and reject late results from an older attempt."),
            ("Job outcome", "Final aggregate and a client-delivery flag.",
             "Avoid recomputation and resend a result whose callback was not acknowledged."),
            ("Timing", "Original acceptance timestamp and remaining budget.",
             "Reconstruct the deadline, with a minimum recovery allowance."),
            ("Ordering metadata", "Per-job sequence, leadership term and coordinator ID.",
             "Fold updates idempotently and reject stale or duplicate state."),
        ],
        widths=[1.05, 3.1, 3.0],
        font_size=9.5,
    )

    heading(doc, "Replicated checkpoint log", level=2)
    body(doc, "Each checkpoint entry has a job ID, a per-job sequence, the producing term and "
              "coordinator ID, and one of three payloads: JobAccepted, SliceUpdated or JobFinished. "
              "Per-job sequence numbers make folding idempotent: a duplicate or out-of-order entry "
              "is ignored instead of applied twice. Every outbound push carries a cumulative "
              "snapshot, so an older asynchronous RPC cannot overwrite a newer view. A JobFinished "
              "entry is a tombstone: it carries the final result (not one slice's result) and a "
              "delivered flag, so a recovered job is not rebuilt and a lost client callback can be "
              "replayed. Finished state has a bounded retention window so a long-lived replica does "
              "not grow forever.")
    body(doc, "The checkpoint log is in memory on every worker. Acceptance is synchronous with one "
              "peer acknowledgement; later slice, completion and delivery updates are asynchronous. "
              "That is a deliberate trade-off between latency and durability. The system is not a "
              "disk-backed Raft/Paxos log: if every replica is lost, or the whole cluster is "
              "restarted, the checkpoints are lost. The report states this limitation rather than "
              "claiming exactly-once durability.")

    heading(doc, "Takeover", level=2)
    body(doc, "When a worker becomes coordinator, it temporarily refuses new jobs. It fetches "
              "checkpoint snapshots from every reachable peer and merges them per job, keeping the "
              "newest sequence for each job. It does not steal a job that is owned by a live "
              "coordinator. For a job whose owner failed, it preserves completed slices and resumes "
              "unfinished slices with fresh discovery. The restored job keeps its original deadline, "
              "subject to a minimum recovery window so a job that was nearly out of time still gets "
              "a fair chance. If every slice is already complete, the job aggregates immediately "
              "instead of waiting for a timeout. If the job had already finished, the saved aggregate "
              "is delivered; delivery is recorded only after the client callback acknowledges it.")
    body(doc, "Recovery has two explicit limits. First, it needs a surviving replica and a reachable "
              "replacement coordinator; a two-worker cluster cannot recover from the coordinator's "
              "failure because there is no majority left. Second, callback delivery is at-least-once "
              "in the presence of lost acknowledgements, so the client deduplicates by job ID. These "
              "limits are documented in docs/failover.md and are visible in the process-level "
              "GrpcOnlyClusterTest, which terminates a real coordinator and checks the recovered "
              "answer and completed progress.")

    heading(doc, "Bootstrap state and membership", level=2)
    body(doc, "The bootstrap is still a single directory with no failover replica, but it no longer "
              "retains crashed workers forever. Workers refresh their membership entry every five "
              "seconds, and the bootstrap reaps an entry after fifteen seconds of silence. That "
              "keeps the majority denominator honest. Worker and job ID counters are persisted in "
              "bootstrap-state.properties (override with --state-file) so a bootstrap restart does "
              "not immediately reuse an identity. Active membership is not persisted: after a "
              "bootstrap restart, workers must re-register, and the documented safe procedure is to "
              "restart the whole cluster if the old membership is gone.")

    # ------------------------------------------------------------------ Q5
    page_break(doc)
    heading(doc, "Q5. Overall failure detection and recovery strategies", level=1)
    body(doc, "DistriLab combines five mechanisms because each addresses a different failure: "
              "unreachability, conflicting leadership, lost work, stale messages and a failed "
              "client callback. The design targets a small, unreliable, unstructured worker network "
              "and must recover without manual restart.")

    heading(doc, "Selected strategies and justification", level=2)
    lead(doc, "Heartbeat monitoring.", "Followers probe the coordinator every 1,000 ms with a 600 ms "
         "deadline. A successful response under 400 ms is ALIVE; a slower response or one or two "
         "misses is SUSPECT; three misses or eight seconds without success is FAILED. This avoids an "
         "election after a single delayed reply, while still detecting a process that has stopped.")
    lead(doc, "Echo-wave election.", "Election messages follow the current neighbour graph and echoes "
         "merge each subtree's best candidate and voters. The lowest JAC wins, with the highest worker "
         "ID as tie-breaker. The initiator announces only with a strict majority of ready workers, "
         "which suits an unstructured overlay better than an ID-ordered Bully or ring.")
    lead(doc, "Majority-quorum protection.", "QuorumMonitor probes ready workers every second, uses a "
         "600 ms deadline, clears the view after two failed checks and renews a 2,500 ms lease. New "
         "jobs require a fresh lease. A leader in a minority partition therefore stops accepting work "
         "instead of creating conflicting results, prioritising safety over availability.")
    lead(doc, "Replicated checkpoints.", "JobAccepted, SliceUpdated and JobFinished entries are "
         "replicated to peers with per-job sequence numbers and terms. At least one peer must confirm "
         "acceptance before the client is told the job was accepted. Later updates are asynchronous, "
         "which keeps submission responsive but means the newest intermediate update may be lost.")
    lead(doc, "Task reassignment and duplicate control.", "When an owner is unreachable, its unfinished "
         "slices are reassigned with an incremented attempt. Stable task IDs and checkpointed results "
         "preserve completed units and reject a late result from an older attempt. This is cheaper "
         "than restarting the whole job and is safer than accepting the first reply blindly.")

    heading(doc, "Alternatives considered", level=2)
    add_table(
        doc,
        ["Alternative", "Attraction", "Why the chosen design differs"],
        [
            ("Fixed central coordinator", "No election and simple testing.",
             "It is a single point of failure and fails the leader-election requirement. The bootstrap remains only a directory."),
            ("Bully or ring election", "Deterministic and low message count on an ordered ring.",
             "An unstructured, healing mesh has no stable ring, and ID order ignores the JAC load metric."),
            ("Gossip membership", "Scales and tolerates churn.",
             "Convergence is slower for a demo and still needs a majority rule to avoid split brain."),
            ("Full Raft or Paxos", "Durable, strongly consistent replicated log.",
             "Much greater implementation and test complexity; the project uses a majority lease and replicated checkpoints, not a consensus library."),
            ("Polling for progress", "No streaming API.",
             "It is not real-time and causes stale bars or heavy polling; gRPC server streaming fits the requirement."),
        ],
        widths=[1.35, 2.15, 3.65],
        font_size=9.5,
    )

    heading(doc, "Recovery sequence", level=2)
    numbered(doc, [
        "Detect an unresponsive worker or coordinator through gRPC deadlines, refused connections "
        "and the health state machine.",
        "For a coordinator failure, start the echo wave and require strict-majority participation "
        "before announcing a replacement.",
        "Install the new coordinator with a higher term; reject stale announcements and step down a "
        "leader that loses its quorum lease.",
        "Fetch checkpoint snapshots from reachable peers, merge by job and sequence, and preserve "
        "completed slices.",
        "Reassign unfinished slices from unreachable owners with new attempts; resume progress "
        "reporting from recovered state.",
        "Redeliver a completed result whose client callback was not acknowledged. Delivery is "
        "at-least-once, so the client deduplicates by job ID.",
    ])

    heading(doc, "Limitations and conclusion", level=2)
    body(doc, "A timeout proves unreachability, not a crash: a slow or isolated process can look "
              "failed. Asynchronous replication can omit the last intermediate checkpoint, and a "
              "cluster without a majority deliberately rejects new work. Checkpoints are in memory, "
              "so losing every replica loses recovery state. The bootstrap is a single directory, "
              "and callbacks are at-least-once. These limits are documented in docs/failover.md and "
              "are preferable to claiming durability the implementation does not provide.")
    body(doc, "The combination is appropriate for the assignment: it elects on an arbitrary graph, "
              "prevents a minority leader from serving, preserves accepted work across a coordinator "
              "change and recovers unfinished slices without restarting every job. Full Raft would "
              "give stronger durability but would replace the assignment's focus with a consensus "
              "implementation. Bounded retention for jobs, tasks, election waves and finished "
              "checkpoints keeps the in-memory design stable during a long demo.")

    # ------------------------------------------------------------------ references
    page_break(doc)
    heading(doc, "References", level=1)
    reference(doc, "Google. (n.d.). gRPC documentation. https://grpc.io/docs/")
    reference(doc, "Google. (n.d.). Protocol Buffers documentation. https://protobuf.dev/")
    reference(doc, "Oracle. (n.d.). Java Remote Method Invocation (Java RMI). "
                   "https://docs.oracle.com/javase/8/docs/technotes/guides/rmi/")
    reference(doc, "Van Steen, M., & Tanenbaum, A. S. (2025). Distributed systems (4th ed., "
                   "Version 4.03). Sections 4.2 and 5.4; Chapters 7 and 8.")

    # ------------------------------------------------------------------ appendix A
    page_break(doc)
    heading(doc, "Appendix A. Git history", level=1)
    body(doc, "Generated from the repository with "
              "git log --graph --date=short --pretty=format:'%h %ad %an%x09%s'. The initial commit "
              "imports the Assignment 1 code, and later commits add the gRPC migration, failover, "
              "quorum and the Assignment 2 hardening.")
    if GIT_LOG.exists():
        code_block(doc, GIT_LOG.read_text(encoding="utf-8"), size=7.5)
    else:
        body(doc, "Git log file was not found. Run the command shown above from the repository root.")

    # ------------------------------------------------------------------ appendix B
    page_break(doc)
    heading(doc, "Appendix B. AI tools and disclosure", level=1)
    body(doc, "AI coding assistants were used during development and for report preparation. They "
              "were treated as tools, not as group members: the group owns the design decisions, "
              "reviewed the output, and is responsible for the submitted code. No AI-generated "
              "change was accepted without a build and test run.")
    add_table(
        doc,
        ["Area", "Tool / mode", "How it was used", "Human verification"],
        [
            ("Election and gRPC migration", "OpenAI Codex via the opencode agent",
             "Read the existing code, proposed the transport seams, wrote/edited classes and tests, ran Maven.",
             "Member reviewed the protocol rules, constants and diffs; green build required."),
            ("Coordinator failover", "OpenAI Codex via the opencode agent",
             "Implemented checkpoint replication, health monitoring and takeover; wrote end-to-end tests and docs.",
             "Member specified the safety rules and reviewed every merge."),
            ("A2 hardening and report", "DeepSeek V4.1 Flash via the opencode agent",
             "Audited compliance, added bootstrap leases/persistent counters/bounded retention/channel pooling, CI, coverage and this Word report.",
             "Changes were compiled and the full suite was run; placeholders and contribution data are filled by the group."),
            ("Documentation", "AI-assisted drafting",
             "Drafted README, protocol notes, failure model and report text from the implementation.",
             "Members checked technical claims against code and test results."),
        ],
        widths=[1.2, 1.55, 2.5, 2.05],
        font_size=9.5,
    )
    body(doc, "Prompting approach: one detailed work order per area, with the frozen-protobuf rule, "
              "required tests and the priority order. The assistant was asked to stop and ask when a "
              "decision affected group ownership. Existing authorship and commit history were "
              "preserved. The historical logs in docs/ai-assisted-audit.md and "
              "docs/ai-usage-member2.md contain more detail, including the bugs found by the "
              "AI-generated tests and the human review decisions.")

    # ------------------------------------------------------------------ appendix C
    page_break(doc)
    heading(doc, "Appendix C. Meeting minutes", level=1)
    body(doc, "The dates, attendees and action owners below are highlighted placeholders. Replace "
              "them with the group's actual minutes before submission.")
    meetings = [
        ("Meeting 1", "[Date]", "[Attendees]",
         "Reviewed Assignment 1 feedback, mapped it to the early refactor branches, split A2 into election, failover, job/progress and bootstrap/test work.",
         "Create the shared proto contracts; agree on the five-job term and JAC rule; set up the module build."),
        ("Meeting 2", "[Date]", "[Attendees]",
         "Reviewed the gRPC migration and the first election tests; agreed that no RMI fallback would remain in production code.",
         "Finish the echo-wave duplicate handling; add chain/triangle tests; document the transport."),
        ("Meeting 3", "[Date]", "[Attendees]",
         "Walked through coordinator failure: heartbeat thresholds, majority lease, acceptance replication and takeover merge rules.",
         "Implement checkpoint replication and the process-level failure test; record known limits."),
        ("Meeting 4", "[Date]", "[Attendees]",
         "Final review of the progress stream, worker reassignment, bootstrap leases and the report package; assigned contribution percentages.",
         "Run mvn verify, fill group details, check the Git log appendix and submit."),
    ]
    for name, date, attendees, decisions, actions in meetings:
        heading(doc, f"{name} — {date}", level=2)
        body(doc, f"Attendees: {attendees}", space_after=3)
        body(doc, f"Decisions: {decisions}", space_after=3)
        body(doc, f"Actions: {actions}", space_after=6)

    # ------------------------------------------------------------------ appendix D
    page_break(doc)
    heading(doc, "Appendix D. Assignment 1 feedback reconciliation", level=1)
    body(doc, "The repository does not store the verbal feedback from the Assignment 1 demo. The "
              "table below reconstructs the feedback themes from the early commits and branches "
              "named in the Git history. Confirm the wording with the group's own notes if a "
              "different list was given.")
    add_table(
        doc,
        ["A1 feedback theme", "Evidence in Git history", "How A2 addresses it"],
        [
            ("Automatically assign worker IDs and ports",
             "e097c6e; refactor/worker-auto-assign-rewire",
             "Bootstrap allocates monotonic IDs; workers bind an available gRPC port and publish it."),
            ("Repair the topology after a partition",
             "e097c6e; refactor/worker-auto-assign-rewire",
             "Failed hops prune edges; repair tries remembered two-hop peers, then discovered workers."),
            ("Show cluster activity in the client",
             "f17f447; refactor/cluster-activity-events",
             "WorkerControlService.GetStatus returns typed recent events; the client has a cluster activity card."),
            ("Warn on very wide PRIMESUM ranges",
             "9163995; refactor/range-warning-1b",
             "RangeWarning warns but allows ranges up to one billion; tests cover the boundary."),
            ("Stream real progress instead of a fake bar",
             "8d1a510; origin/progress-bar",
             "JobProgressService.WatchJob streams completed/total units; clicking a bar opens the worker breakdown."),
            ("Improve client history and scrolling",
             "7fbc9f5; origin/progress-bar",
             "Job history, log scrolling and view actions were added to the client."),
            ("Recover from a worker failure",
             "881a8c0; origin/worker-failure-recovery",
             "Stable task IDs and attempts; unreachable slices are reassigned and completed units are retained."),
            ("Use a launcher instead of a long classpath command",
             "8a6ff4a; refactor/run-script",
             "run.ps1 builds and starts bootstrap, workers and clients; the classpath is documented."),
            ("Freeze the gRPC contracts",
             "aa7ce21; feature/grpc-contracts",
             "Six proto files with frozen field numbers; ProtoContractsTest checks the shapes."),
            ("Elect a coordinator and survive its failure",
             "leader-election, coordinator-failover, feature/grpc-majority-quorum",
             "Echo-wave election, strict-majority lease, checkpoint replication and takeover."),
        ],
        widths=[1.65, 2.1, 3.35],
        font_size=9.5,
    )

    # ------------------------------------------------------------------ appendix E
    page_break(doc)
    heading(doc, "Appendix E. Verification and reproducibility", level=1)
    body(doc, "The complete reactor build and test suite was run from the repository root with JDK "
              "25 and Maven 3.9.x:")
    code_block(doc, "mvn -B verify", size=10)
    body(doc, "Result: BUILD SUCCESS, 179 tests run, 0 failures, 0 errors, 0 skipped. The suite "
              "includes actual worker JVMs, a forcibly terminated coordinator, stream reconnects, "
              "bootstrap lease expiry, persistent counter restart, bounded retention, the "
              "multi-worker splitter scale check and the no-RMI guard. JaCoCo is configured in the "
              "root POM, and integration-tests/pom.xml produces the aggregate report because all "
              "tests live in that module. The GitHub Actions workflow "
              ".github/workflows/verify.yml runs the same command on JDK 25 and uploads Surefire "
              "and coverage reports.")
    body(doc, "Setup and run instructions are in README.md. The normal demo uses one bootstrap, "
              "four workers and one client. Progress is visible in the client; stopping a worker "
              "shows reassignment, and stopping the coordinator shows election and recovery. The "
              "bootstrap state file is git-ignored and should not be committed.")

    highlight_placeholders(doc)
    doc.core_properties.title = "CS324 Assignment 2 - gRPC Worker Cluster"
    doc.core_properties.author = "CS324 group"
    doc.core_properties.subject = "Leader election, failure tolerance and progress streaming"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
